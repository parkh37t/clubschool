#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Claude-native md → DOCX 변환기 (Adobe 불필요).
PROJECT/*.md (SSOT) → DELIVERABLES/.../*.docx. 정책: Claude-native Office First.
헤딩/불릿/번호/표/인용을 IBK 브랜드 스타일로 변환.
"""
import os, re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

KO = "Malgun Gothic"
BLUE = RGBColor(0x0A, 0x4D, 0xA2)
FG = RGBColor(0x1A, 0x1F, 0x2B)
MUT = RGBColor(0x5A, 0x64, 0x73)
BASE = os.path.dirname(os.path.abspath(__file__))
PROOT = os.path.join(BASE, "..", "..")

def _clean(s):  # 인라인 마크다운 제거
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    s = re.sub(r"`(.+?)`", r"\1", s)
    s = s.replace("**", "").replace("`", "")
    return s.strip()

def _runf(run, size=10.5, color=FG, bold=False):
    run.font.name = KO; run.font.size = Pt(size); run.font.color.rgb = color; run.font.bold = bold

def _is_table_row(l): return l.strip().startswith("|") and l.strip().endswith("|")
def _is_sep(l): return bool(re.match(r"^\s*\|?[\s:|-]+\|?\s*$", l)) and "-" in l

def add_table(doc, rows):
    cells = [[_clean(c) for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [c for c in cells if not all(re.match(r"^[\s:-]*$", x) for x in c)]
    if not cells: return
    ncol = max(len(r) for r in cells)
    t = doc.add_table(rows=0, cols=ncol); t.style = "Light Grid Accent 1"
    for ri, row in enumerate(cells):
        rcells = t.add_row().cells
        for ci in range(ncol):
            val = row[ci] if ci < len(row) else ""
            p = rcells[ci].paragraphs[0]; r = p.add_run(val)
            _runf(r, 9, BLUE if ri == 0 else FG, bold=(ri == 0))

def md_to_docx(md_paths, out_path, doc_title):
    doc = Document()
    doc.styles["Normal"].font.name = KO; doc.styles["Normal"].font.size = Pt(10.5)
    tp = doc.add_paragraph(); tr = tp.add_run(doc_title); _runf(tr, 20, BLUE, bold=True)
    sub = doc.add_paragraph(); _runf(sub.add_run("Claude-native DOCX · SSOT: PROJECT/*.md"), 9, MUT)

    for idx, md_path in enumerate(md_paths):
        if idx > 0:
            doc.add_page_break()
        with open(md_path, encoding="utf-8") as f:
            lines = f.read().splitlines()
        i = 0
        while i < len(lines):
            l = lines[i]
            if _is_table_row(l):
                blk = []
                while i < len(lines) and _is_table_row(lines[i]):
                    blk.append(lines[i]); i += 1
                add_table(doc, blk); continue
            s = l.strip()
            if not s:
                i += 1; continue
            if s.startswith("```"):  # code fence: skip fence, keep content mono
                i += 1
                while i < len(lines) and not lines[i].strip().startswith("```"):
                    p = doc.add_paragraph(); r = p.add_run(lines[i]); r.font.name = "Consolas"; r.font.size = Pt(9); r.font.color.rgb = MUT
                    i += 1
                i += 1; continue
            m = re.match(r"^(#{1,4})\s+(.*)", s)
            if m:
                lvl = len(m.group(1)); txt = _clean(m.group(2))
                p = doc.add_paragraph(); r = p.add_run(txt)
                _runf(r, {1:17,2:14,3:12,4:11}[lvl], BLUE if lvl <= 2 else FG, bold=True)
                p.space_before = Pt(8 if lvl <= 2 else 4); p.space_after = Pt(3)
            elif s.startswith(("> ",)):
                p = doc.add_paragraph(); r = p.add_run(_clean(s[2:])); _runf(r, 9.5, MUT); r.italic = True
            elif re.match(r"^[-*•]\s+", s):
                p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.25)
                _runf(p.add_run("• " + _clean(re.sub(r"^[-*•]\s+", "", s))), 10.5, FG)
            elif re.match(r"^\d+\.\s+", s):
                p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.2)
                _runf(p.add_run(_clean(s)), 10.5, FG)
            else:
                _runf(doc.add_paragraph().add_run(_clean(s)), 10.5, FG)
            i += 1
    doc.save(out_path); return out_path

if __name__ == "__main__":
    # 과제 무관 CLI 엔진: 특정 프로젝트에 고정하지 않는다.
    # 사용법: gen_docs_from_md.py <out.docx> "<title>" <md1> [md2 ...]
    # 프로젝트별 일괄 생성은 _generators/projects/<slug>.sh 가 이 엔진을 호출한다.
    import sys
    args = sys.argv[1:]
    if len(args) < 3:
        print('usage: gen_docs_from_md.py <out.docx> "<title>" <md1> [md2 ...]')
        sys.exit(1)
    out, title, mds = args[0], args[1], args[2:]
    os.makedirs(os.path.dirname(out) or ".", exist_ok=True)
    md_to_docx(mds, out, title)
    print("OK", os.path.getsize(out), out)
